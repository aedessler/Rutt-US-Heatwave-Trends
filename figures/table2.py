#!/usr/bin/env python3
"""Table 2 -- CONUS JJA temperature anomalies, the numbers Figure 2 drew as bars.

Same anomaly-first pipeline and same QC as ``figure2.py``, but the record runs
to 2026 rather than 2025, GHCN-Daily and the USHCN-masked Berkeley leg are gone
(Figure 2 computed both and plotted neither), and the summary is a table rather
than a bar chart:

    1934 | 1936 | 2024 | 2026 | 1930-1939 mean | 2017-2026 mean

Datasets whose record stops before 2026 get the mean of 2017 through their own
last JJA instead, and that end year is reported in the table.

The gridded station fields are cached under the SAME key Figure 1 uses
(``grid_<ds>_<elem>_p2_k8_r150_1900_2026``): identical build code, identical
gridder, identical span, so the table reuses Figure 1's caches rather than
regridding.

    python figures/table2.py
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import *          # noqa: F401,F403

check_paths()
print()
check_geometry()
print()
check_pipeline()
print()
check_gridded()
print()

warnings.filterwarnings("ignore", category=RuntimeWarning)
CACHE_ANOM.mkdir(exist_ok=True)
FIGD.mkdir(exist_ok=True)

# ══ SETTINGS ════════════════════════════════════════════════════════════════
YEARS_T   = np.arange(1900, 2027)      # one year past Figure 2: 2026 JJA is in
SEASON_T  = (6, 7, 8)                  # JJA

MIN_COV_FRAC_T   = 0.50    # season dropped if less than this share of CONUS reported
MIN_PERIOD_YRS_T = 7       # of the 10 years in a period mean
MAD_K_T = 8.0
IDW_T = dict(power=2.0, k=8, radius_km=150.0)

SINGLE_YEARS_T = (1934, 1936, 2024, 2026)
DUST_START_T, DUST_END_T     = 1930, 1939
MODERN_START_T, MODERN_END_T = 2017, 2026

UH_LABEL_RAW_T  = "USHCN-Daily"
UH_LABEL_ADJ_T  = "USHCN-BC"

# Figure 2 computed GHCN-Daily and the USHCN-masked Berkeley leg ("BE @ USHCN")
# but suppressed both from the plot; the table drops them outright.
LABELS_T = {"berkeley": "Berkeley", "nclimdiv": "nCLIMDIV",
            "noaa": "NOAAGlobalTemp", "crutem5": "CRUTEM5",
            "ushcn": UH_LABEL_RAW_T, "ushcn_bc": UH_LABEL_ADJ_T}
ORDER_T = {"tmax": ["berkeley", "nclimdiv", "ushcn", "ushcn_bc"],
           "tmin": ["berkeley", "nclimdiv", "ushcn", "ushcn_bc"],
           "tavg": ["berkeley", "nclimdiv", "noaa", "crutem5", "ushcn", "ushcn_bc"]}
ELEM_TITLES_T = {"tmax": "TMAX  -  daily maximum temperature",
                 "tmin": "TMIN  -  daily minimum temperature",
                 "tavg": "TAVG  -  daily mean temperature"}

# ══ QC AND SERIES BUILDERS (Figure 2's, verbatim apart from the year span) ══
def _mad_filt_T(s, k=MAD_K_T):
    v = np.asarray(s.values, float)
    med = np.nanmedian(v); mad = np.nanmedian(np.abs(v-med))
    if not np.isfinite(mad) or mad == 0: return s
    bad = np.isfinite(v) & (np.abs(v-med) > k*1.4826*mad)
    return s.where(~pd.Series(bad, index=s.index)) if bad.any() else s

def _trim_tail_T(s, z_k=6.0):
    v = np.asarray(s.values, float); fi = np.where(np.isfinite(v))[0]
    if len(fi) < 15: return s
    hist, tail = fi[:-2], fi[-2:]
    if len(hist) < 10: return s
    med = np.nanmedian(v[hist]); mad = np.nanmedian(np.abs(v[hist]-med))
    if not np.isfinite(mad) or mad == 0: return s
    bad = [i for i in tail if np.isfinite(v[i]) and abs(v[i]-med) > z_k*1.4826*mad]
    if not bad: return s
    keep = np.ones_like(v, bool); keep[bad] = False
    return s.where(pd.Series(keep, index=s.index))

def seasonal_from_field_T(field, weights, gate_coverage=False):
    """cell anomalies -> area mean -> day-weighted JJA -> QC -> Series by year."""
    m, frac = area_mean(cell_anomalies(field), weights)
    s = pd.Series(season_mean(_to_year_month(m, YEARS_T), SEASON_T, YEARS_T), index=YEARS_T)
    if gate_coverage:
        f = pd.Series(season_mean(_to_year_month(frac, YEARS_T), SEASON_T, YEARS_T),
                      index=YEARS_T)
        s = s.where(f >= MIN_COV_FRAC_T)
    return _trim_tail_T(_mad_filt_T(s))

def station_field_T(dataset, elem, gridder):
    # Figure 1's cache key: same build, same span, so the cache is shared.
    key = (f"grid_{dataset}_{elem}_p{IDW_T['power']:g}_k{IDW_T['k']}"
           f"_r{IDW_T['radius_km']:g}_{YEARS_T[0]}_{YEARS_T[-1]}")
    def build():
        sm = station_anomalies(station_months(dataset, elem, YEARS_T))
        return grid_station_field(sm, gridder, YEARS_T)
    return _cache(key, build)

def berkeley_field_T(fp, elem):
    """All 12 months of monthly means; Figure 2's ``field16_berkeley_*`` cache."""
    def build():
        ds = xr.open_dataset(str(fp))
        t = _std(ds["temperature"]).load(); ds.close()
        u = (t.attrs.get("units", "") or "").lower()
        if u in ["k", "kelvin", "degk"] or "kelvin" in u: t = t - 273.15
        m = t.resample(time="MS").mean()
        return m.where(t.resample(time="MS").count() >= MIN_DAYS_MON)
    return _cache(f"field16_berkeley_{elem}", build)

# nCLIMDIV: NOAA's OFFICIAL area-weighted CONUS national mean (region 110).
_NCLIMDIV_FILES_T = {e: nclimdiv_file(e) for e in ("tmax", "tmin", "tavg")}
_NCLIMDIV_ELEM_T  = {"tmax": 27, "tmin": 28, "tavg": 2}

def nclimdiv_seasonal_T(elem):
    fp = _NCLIMDIV_FILES_T[elem]
    prefix = f"1100{_NCLIMDIV_ELEM_T[elem]:02d}"
    ym = np.full((len(YEARS_T), 12), np.nan)
    with open(str(fp)) as f:
        for line in f:
            if line[:6] == prefix:
                yr = int(line[6:10])
                if not (YEARS_T[0] <= yr <= YEARS_T[-1]): continue
                for mo in range(12):
                    v = float(line[10+mo*7:17+mo*7])
                    if v > -90.0: ym[yr-YEARS_T[0], mo] = (v-32.0)*5.0/9.0
    base = ym[(YEARS_T >= BASELINE[0]) & (YEARS_T <= BASELINE[1])]
    clim = np.where(np.isfinite(base).sum(0) >= MIN_BASE_CELL, np.nanmean(base, 0), np.nan)
    s = pd.Series(season_mean(ym-clim[None, :], SEASON_T, YEARS_T), index=YEARS_T)
    return _trim_tail_T(_mad_filt_T(s))

# ══ BUILD EVERY SERIES ══════════════════════════════════════════════════════
_lat05_T, _lon05_T = g05_centers()
W05_T  = cell_weights(_lat05_T, _lon05_T, key="g05")
GRID_T = IDWGrid(_lat05_T, _lon05_T, **IDW_T)

SER_T = {}                                   # (dataset, elem) -> Series indexed by year
for _ds in ("ushcn", "ushcn_bc"):
    for _e in ("tmax", "tmin", "tavg"):
        SER_T[(_ds, _e)] = seasonal_from_field_T(station_field_T(_ds, _e, GRID_T), W05_T)

_be_f_T = {"tmax": berkeley_field_T(BE_TMAX, "tmax"),
           "tmin": berkeley_field_T(BE_TMIN, "tmin")}
if BE_TAVG.exists():
    _be_f_T["tavg"] = berkeley_field_T(BE_TAVG, "tavg")
_Wbe_T = cell_weights(_be_f_T["tmax"]["lat"].values, _be_f_T["tmax"]["lon"].values,
                      key="berkeley")
for _e, _f in _be_f_T.items():
    SER_T[("berkeley", _e)] = seasonal_from_field_T(_f, _Wbe_T)
if "tavg" not in _be_f_T:                     # no Berkeley TAVG file: 0.5*(TMAX+TMIN)
    SER_T[("berkeley", "tavg")] = 0.5*(SER_T[("berkeley", "tmax")]
                                       + SER_T[("berkeley", "tmin")])

_noaa_T = noaa_field()
SER_T[("noaa", "tavg")] = seasonal_from_field_T(
    _noaa_T, cell_weights(_noaa_T["lat"].values, _noaa_T["lon"].values, key="noaa"),
    gate_coverage=True)
_cru_T = crutem_field()
SER_T[("crutem5", "tavg")] = seasonal_from_field_T(
    _cru_T, cell_weights(_cru_T["lat"].values, _cru_T["lon"].values, key="crutem5"),
    gate_coverage=True)
for _e in ("tmax", "tmin", "tavg"):
    SER_T[("nclimdiv", _e)] = nclimdiv_seasonal_T(_e)

# ══ THE SIX COLUMNS ═════════════════════════════════════════════════════════
def _yr_T(s, year):
    v = s.get(year, np.nan)
    return float(v) if np.isfinite(v) else np.nan

def _pm_T(s, y0, y1, min_yrs=MIN_PERIOD_YRS_T):
    v = np.asarray(s.loc[y0:y1].values, float)
    return float(np.nanmean(v)) if np.isfinite(v).sum() >= min_yrs else np.nan

def _last_year_T(s):
    lv = s.last_valid_index()
    return None if lv is None else int(min(lv, MODERN_END_T))

def row_T(elem, key):
    """One dataset's row: the four single years, the Dust Bowl mean and the
    modern mean over 2017 through the dataset's own last JJA."""
    s = SER_T.get((key, elem))
    if s is None: return None
    end = _last_year_T(s)
    if end is None: return None
    rec = {"element": elem.upper(), "dataset": LABELS_T[key], "key": key}
    for y in SINGLE_YEARS_T:
        rec[str(y)] = _yr_T(s, y)
    rec[f"{DUST_START_T}-{DUST_END_T}"] = _pm_T(s, DUST_START_T, DUST_END_T)
    rec["modern_start"] = MODERN_START_T
    rec["modern_end"]   = end
    rec["modern_mean"]  = _pm_T(s, MODERN_START_T, end,
                                min_yrs=min(MIN_PERIOD_YRS_T, end-MODERN_START_T+1))
    rec["last_year"]    = int(s.last_valid_index())
    return rec

TAB_T = pd.DataFrame([r for e in ("tmax", "tmin", "tavg") for k in ORDER_T[e]
                      for r in [row_T(e, k)] if r is not None])
TAB_T["modern_window"] = [f"{a}-{b}" for a, b in zip(TAB_T.modern_start, TAB_T.modern_end)]

# ══ RENDER ══════════════════════════════════════════════════════════════════
_DUST_COL_T = f"{DUST_START_T}-{DUST_END_T}"
_VAL_COLS_T = [str(y) for y in SINGLE_YEARS_T] + [_DUST_COL_T, "modern_mean"]
_HEADS_T    = [str(y) for y in SINGLE_YEARS_T] + [_DUST_COL_T, f"{MODERN_START_T}-{MODERN_END_T}"]

def _f_T(v):
    return "  n/a" if v is None or not np.isfinite(v) else f"{v:+.2f}"

_W_T = [12]*(len(_HEADS_T)-1) + [20]   # last column carries a "[2017-end]" tag

_lines_T = []
_lines_T.append("Table 2. CONUS JJA temperature anomalies (deg C vs. 1951-1980).")
_lines_T.append("Modern-mean column is 2017-2026 where the record allows; a dataset that")
_lines_T.append("stops earlier is averaged 2017 through its own last JJA, shown in brackets.")
for _e in ("tmax", "tmin", "tavg"):
    sub = TAB_T[TAB_T.element == _e.upper()]
    _lines_T.append("")
    _lines_T.append(ELEM_TITLES_T[_e])
    _lines_T.append(f"{'dataset':<16}" + "".join(f"{h:>{w}}" for h, w in zip(_HEADS_T, _W_T)))
    for _, r in sub.iterrows():
        cells = [_f_T(r[c]) for c in _VAL_COLS_T]
        if r.modern_end != MODERN_END_T:
            cells[-1] = f"{cells[-1]} [2017-{r.modern_end}]"
        _lines_T.append(f"{r.dataset:<16}" + "".join(f"{c:>{w}}" for c, w in zip(cells, _W_T)))
_txt_T = "\n".join(_lines_T)
print(_txt_T)

# Markdown and CSV are off for now -- the Word table is the deliverable.
# Uncomment this block and the two writes at the bottom to get them back.
# # ── markdown, for dropping straight into the manuscript ──────────────────────
# _md_T = ["# Table 2. CONUS JJA temperature anomalies (°C, vs. 1951–1980)", "",
#          "Modern mean is 2017–2026 where the record allows; a dataset that stops",
#          "earlier is averaged 2017 through its own last JJA (given in brackets).", ""]
# for _e in ("tmax", "tmin", "tavg"):
#     sub = TAB_T[TAB_T.element == _e.upper()]
#     _md_T += [f"## {ELEM_TITLES_T[_e].replace('  -  ', ' — ')}", "",
#               "| Dataset | " + " | ".join(h.replace("-", "–") for h in _HEADS_T) + " |",
#               "|" + "---|"*(len(_HEADS_T)+1)]
#     for _, r in sub.iterrows():
#         cells = [("n/a" if not np.isfinite(r[c]) else f"{r[c]:+.2f}") for c in _VAL_COLS_T]
#         if r.modern_end != MODERN_END_T:
#             cells[-1] = f"{cells[-1]} [2017–{r.modern_end}]"
#         _md_T.append(f"| {r.dataset} | " + " | ".join(cells) + " |")
#     _md_T.append("")

# ── Word, for the manuscript ────────────────────────────────────────────────
# Journal house style: no vertical rules and no interior grid, a heavy rule
# above and below the table, a light one under the column heads and above each
# element block. Datasets whose record stops before 2026 carry a superscript
# marker on the modern-mean cell rather than an inline "[2017-2024]", which
# would widen the column for the sake of three rows.
DOCX_FONT_T, DOCX_PT_T, DOCX_HEAD_PT_T = "Times New Roman", 9.0, 9.0
DOCX_COL_IN_T = [1.45] + [0.72]*4 + [0.86, 0.86]   # dataset + 4 years + 2 means
DOCX_RULE_HEAVY_T, DOCX_RULE_LIGHT_T = 12, 6       # eighths of a point
DOCX_BLOCK_SHADE_T = "F2F2F2"                      # element sub-header fill
NA_T = "—"                                    # em dash for "no such season"


def _tc_borders_T(cell, **edges):
    """Per-cell borders: _tc_borders_T(cell, top=12, bottom=6). Anything not
    named is set to none, so cells start clean and only the rules we ask for
    are drawn."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    el = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        sz = edges.get(edge)
        e.set(qn("w:val"), "single" if sz else "none")
        e.set(qn("w:sz"), str(sz or 0))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        el.append(e)
    tcPr.append(el)


def _shade_T(cell, hexfill):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")          # never "solid": it renders black
    sh.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(sh)


def _put_T(cell, runs, bold=False, align="center", width_in=None):
    """Write one cell. `runs` is a string, or (text, is_superscript) pairs --
    python-docx has no rich-text setter, so the runs are built by hand."""
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    if width_in is not None:
        cell.width = Inches(width_in)
    para = cell.paragraphs[0]
    para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                      "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    pf = para.paragraph_format
    pf.space_before, pf.space_after = Pt(1.5), Pt(1.5)
    for text, sup in ([(runs, False)] if isinstance(runs, str) else runs):
        r = para.add_run(text)
        r.font.name, r.font.size, r.bold = DOCX_FONT_T, Pt(DOCX_PT_T), bold
        r.font.superscript = sup or None


def write_docx_T(fp):
    """Table 2 as a .docx: caption, the three element blocks, and the note."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = DOCX_FONT_T, Pt(DOCX_PT_T)
    # python-docx sets only the latin font; Word also wants the east-asian slot
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT_T)

    # markers for the datasets whose modern window falls short of 2026
    short = sorted({int(r.modern_end) for _, r in TAB_T.iterrows()
                    if r.modern_end != MODERN_END_T})
    marker = {end: chr(ord("a") + i) for i, end in enumerate(short)}

    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(6)
    for text, bold in ((f"Table 2. ", True),
                       ("CONUS JJA temperature anomalies (°C, relative to the "
                        f"{BASELINE[0]}–{BASELINE[1]} mean).", False)):
        r = cap.add_run(text)
        r.font.name, r.font.size, r.bold = DOCX_FONT_T, Pt(DOCX_PT_T), bold

    heads = ["Dataset"] + [h.replace("-", "–") for h in _HEADS_T]
    nrows = 1 + sum(1 + (TAB_T.element == e.upper()).sum() for e in ("tmax", "tmin", "tavg"))
    tbl = doc.add_table(rows=nrows, cols=len(heads))
    tbl.style = doc.styles["Table Grid"]      # a real border set to override
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for col, w in zip(tbl.columns, DOCX_COL_IN_T):
        col.width = Inches(w)

    for c, (cell, w) in enumerate(zip(tbl.rows[0].cells, DOCX_COL_IN_T)):
        _put_T(cell, heads[c], bold=True, align="left" if c == 0 else "center", width_in=w)
        _tc_borders_T(cell, top=DOCX_RULE_HEAVY_T, bottom=DOCX_RULE_LIGHT_T)

    ri = 1
    for bi, e in enumerate(("tmax", "tmin", "tavg")):
        sub = TAB_T[TAB_T.element == e.upper()]
        row = tbl.rows[ri]
        for cell, w in zip(row.cells, DOCX_COL_IN_T):
            _tc_borders_T(cell, top=0 if bi == 0 else DOCX_RULE_LIGHT_T)
            _shade_T(cell, DOCX_BLOCK_SHADE_T)
            cell.width = Inches(w)
        span = row.cells[0].merge(row.cells[-1])
        _put_T(span, ELEM_TITLES_T[e].replace("  -  ", " — "), bold=True, align="left")
        ri += 1
        for _, r in sub.iterrows():
            last = ri == nrows - 1
            for c, (cell, w) in enumerate(zip(tbl.rows[ri].cells, DOCX_COL_IN_T)):
                _tc_borders_T(cell, bottom=DOCX_RULE_HEAVY_T if last else 0)
                if c == 0:
                    _put_T(cell, r.dataset, align="left", width_in=w)
                    continue
                col = _VAL_COLS_T[c-1]
                v = r[col]
                txt = NA_T if not np.isfinite(v) else f"{v:+.2f}"
                runs = [(txt, False)]
                if col == "modern_mean" and r.modern_end != MODERN_END_T:
                    runs.append((marker[int(r.modern_end)], True))
                _put_T(cell, runs, width_in=w)
            ri += 1

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    pieces = [("Note. ", True),
              (f"Anomalies are June–August means relative to {BASELINE[0]}–"
               f"{BASELINE[1]}. {NA_T} marks a season the record does not cover. ", False)]
    for end, mk in marker.items():
        pieces += [(mk, "sup"),
                   (f" mean over {MODERN_START_T}–{end}, the dataset's last "
                    f"complete JJA. ", False)]
    for text, kind in pieces:
        r = note.add_run(text)
        r.font.name, r.font.size = DOCX_FONT_T, Pt(DOCX_PT_T)
        r.bold = kind is True
        r.font.superscript = True if kind == "sup" else None

    doc.save(str(fp))


# _out_csv_T  = FIGD / "Table2.csv"
# _out_md_T   = FIGD / "Table2.md"
_out_docx_T = FIGD / "Table2.docx"
# TAB_T[["element", "dataset", "key"] + [str(y) for y in SINGLE_YEARS_T]
#       + [_DUST_COL_T, "modern_window", "modern_mean", "last_year"]
#       ].round(4).to_csv(_out_csv_T, index=False)
# _out_md_T.write_text("\n".join(_md_T))

# print(f"\nwrote {_out_csv_T}")
# print(f"wrote {_out_md_T}")
print()
try:
    write_docx_T(_out_docx_T)
    print(f"wrote {_out_docx_T}")
except ImportError:
    print(f"skipped {_out_docx_T}: python-docx is not installed (pip install python-docx)")
